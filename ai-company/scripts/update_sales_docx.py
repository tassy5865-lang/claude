#!/usr/bin/env python3
"""営業Wordテンプレートを証拠ベースの短文構造へ更新する。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "sales" / "private" / "営業文.docx"

NAVY = RGBColor(0x0B, 0x25, 0x45)
CYAN = RGBColor(0x08, 0x7E, 0x8B)
MUTED = RGBColor(0x48, 0x65, 0x81)
FONT = "Yu Gothic"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_free_rule(paragraph, color="D9E2EC"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


doc = Document(DOCX)
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

section = doc.sections[0]
section.start_type = WD_SECTION.NEW_PAGE
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.2

title_style = doc.styles["Title"]
title_style.font.name = FONT
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = NAVY
title_style.paragraph_format.space_after = Pt(5)

for style_name, size, before, after in (
    ("Heading 1", 16, 15, 7),
    ("Heading 2", 12.5, 10, 4),
):
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = CYAN if style_name == "Heading 1" else NAVY
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
header_p = header.paragraphs[0]
header_p.text = "Nexsist.cs  |  Sales Template"
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header_p.runs[0], size=8.5, bold=True, color=MUTED)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "INTERNAL TEMPLATE - 連絡先は送信時に人間が入力"
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer_p.runs[0], size=8, color=MUTED)

p = doc.add_paragraph(style="Title")
p.add_run("Nexsist.cs 応募文テンプレート")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("AI・DX・業務改善案件向け / 送信本文 250〜400字")
set_font(r, size=11.5, bold=True, color=MUTED)
set_cell_free_rule(p, "087E8B")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("目的: ")
set_font(r, bold=True, color=CYAN)
r = p.add_run("相手の課題、提供する成果物、確認済みの証拠、次の一歩を短く伝える。長い自己紹介は送信本文に入れない。")
set_font(r)

doc.add_paragraph("1. 送信前ゲート", style="Heading 1")
for label, value in (
    ("案件証拠", "直接URL・取得日時・最終確認日時・募集状態がそろい、応募候補が「可」"),
    ("実績証拠", "sales/private/evidence.md で「確認済み」かつ「応募文への記載可」"),
    ("データ", "顧客名・患者情報・秘密情報を含めない。必要なら匿名化"),
    ("連絡先", "本文では [送信時に人間が入力] のままにする"),
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=NAVY)
    r = p.add_run(value)
    set_font(r)

doc.add_paragraph("2. 送信本文", style="Heading 1")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.18)
p.paragraph_format.right_indent = Inches(0.18)
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(8)
template_text = (
    "【相手の課題を募集文の言葉で1文】を拝見し、【具体的な成果物】でお力になれると考えました。"
    "【確認済み実績を1件だけ。証拠IDも記録】。まず【小さな初回成果物】をご提示し、"
    "操作や出力をご確認いただいてから本実装へ進みます。【データ安全上の要点】。"
    "最初に【返信しやすい確認質問を1つ】を伺えますでしょうか。"
)
r = p.add_run(template_text)
set_font(r, size=11)

doc.add_paragraph("本文に含める4要素", style="Heading 2")
for label, text in (
    ("課題", "相手の困りごと"),
    ("成果", "納品物または改善結果"),
    ("証拠", "確認済みの証拠を1つ"),
    ("質問", "返信しやすい質問を1つ"),
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r)

doc.add_paragraph("3. 作成記録", style="Heading 1")
for text in (
    "案件名: [入力]",
    "文字数: [250〜400字で確認]",
    "使用した証拠ID: [入力 / なければ「なし」]",
    "募集状態の最終確認: [YYYY-MM-DD HH:mm JST]",
    "確認者: [人間が入力]",
    "連絡先: [送信時に人間が入力]",
):
    p = doc.add_paragraph(text)
    p.paragraph_format.left_indent = Inches(0.18)

doc.add_paragraph("4. 使用禁止の表現", style="Heading 1")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.18)
r = p.add_run(
    "根拠のない「現場で運用」「導入済み」「月○時間削減」、未確認の資格・顧客名、"
    "「何でも対応可能」「精一杯頑張ります」など成果が伝わらない表現は使用しない。"
)
set_font(r)

doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.title = "Nexsist.cs 応募文テンプレート"
doc.core_properties.subject = "証拠ベースの短文営業テンプレート"
doc.core_properties.comments = "個人連絡先を含まない社内テンプレート"
doc.save(DOCX)
print(DOCX)
